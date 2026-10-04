"""Tool registry and built-in tools. The AI can call any registered tool.

Tools marked confirm=True change the machine or contact people, so the UI
asks the user to approve them before they run.
"""
import base64, io, os, platform, re, smtplib, sqlite3, subprocess, time, webbrowser
from concurrent.futures import ThreadPoolExecutor
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import quote_plus

import httpx

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
UPLOADS.mkdir(parents=True, exist_ok=True)
OS = platform.system()
TOOLS: dict = {}


def tool(name, desc, params, confirm=False):
    """Register a function as an AI-callable tool. params: {arg: "description"}."""
    props = {k: {"type": "string", "description": v} for k, v in params.items()}
    def deco(fn):
        TOOLS[name] = {"fn": fn, "confirm": confirm, "schema": {"type": "function", "function": {
            "name": name, "description": desc,
            "parameters": {"type": "object", "properties": props, "required": list(props)}}}}
        return fn
    return deco


# ---------- memory + reminders (SQLite) ----------
def db():
    c = sqlite3.connect(DATA / "chemistry_ai.db")
    c.execute("CREATE TABLE IF NOT EXISTS memory(id INTEGER PRIMARY KEY, fact TEXT UNIQUE)")
    c.execute("CREATE TABLE IF NOT EXISTS reminders(id INTEGER PRIMARY KEY, text TEXT, due REAL, done INTEGER DEFAULT 0)")
    return c


def memories() -> list[tuple[int, str]]:
    with db() as c:
        return c.execute("SELECT id, fact FROM memory ORDER BY id").fetchall()


@tool("remember", "Save a fact about the user to long-term memory.", {"fact": "The fact to remember"})
def remember(fact):
    with db() as c:
        c.execute("INSERT OR IGNORE INTO memory(fact) VALUES(?)", (fact.strip()[:300],))
    return "Saved."


@tool("set_reminder", "Set a reminder that fires after some minutes.", {"text": "What to remind", "minutes": "Minutes from now (number)"})
def set_reminder(text, minutes):
    with db() as c:
        c.execute("INSERT INTO reminders(text, due) VALUES(?,?)", (text[:200], time.time() + float(minutes) * 60))
    return f"Reminder set for {minutes} minutes from now."


def due_reminders() -> list[str]:
    with db() as c:
        rows = c.execute("SELECT id, text FROM reminders WHERE done=0 AND due<=?", (time.time(),)).fetchall()
        c.executemany("UPDATE reminders SET done=1 WHERE id=?", [(r[0],) for r in rows])
    return [r[1] for r in rows]


# ---------- system control ----------
@tool("set_volume", "Set the computer's speaker volume.", {"percent": "0-100"}, confirm=True)
def set_volume(percent):
    n = max(0, min(100, int(float(percent))))
    if OS == "Darwin":
        subprocess.run(["osascript", "-e", f"set volume output volume {n}"], check=True)
    elif OS == "Windows":
        from pycaw.pycaw import AudioUtilities
        AudioUtilities.GetSpeakers().EndpointVolume.SetMasterVolumeLevelScalar(n / 100, None)
    else:
        subprocess.run(["pactl", "set-sink-volume", "@DEFAULT_SINK@", f"{n}%"], check=True)
    return f"Volume set to {n}%."


@tool("set_brightness", "Set the screen brightness.", {"percent": "0-100"}, confirm=True)
def set_brightness(percent):
    import screen_brightness_control as sbc
    sbc.set_brightness(max(0, min(100, int(float(percent)))))
    return "Brightness updated."


@tool("open_app", "Open an application by name (e.g. Calculator, Chrome).", {"name": "Application name"}, confirm=True)
def open_app(name):
    if not re.fullmatch(r"[\w .\-]{1,60}", name):
        return "Invalid application name."
    if OS == "Windows":
        subprocess.Popen(["cmd", "/c", "start", "", name])
    elif OS == "Darwin":
        subprocess.Popen(["open", "-a", name])
    else:
        subprocess.Popen([name.lower().replace(" ", "-")])
    return f"Opening {name}."


# ---------- web: search, weather ----------
@tool("web_search", "Search the web and return the top results.", {"query": "Search query"})
def web_search(query):
    from ddgs import DDGS
    rows = DDGS().text(query, max_results=5)
    return "\n".join(f"- {r['title']}: {r['body']} ({r['href']})" for r in rows) or "No results."


@tool("get_weather", "Get the current weather for a city.", {"city": "City name"})
def get_weather(city):
    g = httpx.get("https://geocoding-api.open-meteo.com/v1/search", params={"name": city, "count": 1}, timeout=15).json()
    if not g.get("results"):
        return "City not found."
    p = g["results"][0]
    w = httpx.get("https://api.open-meteo.com/v1/forecast", timeout=15, params={
        "latitude": p["latitude"], "longitude": p["longitude"],
        "current": "temperature_2m,relative_humidity_2m,wind_speed_10m,weather_code"}).json()["current"]
    return f"{p['name']}: {w['temperature_2m']} C, humidity {w['relative_humidity_2m']}%, wind {w['wind_speed_10m']} km/h (WMO code {w['weather_code']})."


# ---------- browser + YouTube (Playwright, one dedicated thread) ----------
_pw = ThreadPoolExecutor(max_workers=1)
_state: dict = {}


def _page():
    if "page" not in _state:
        from playwright.sync_api import sync_playwright
        _state["pw"] = sync_playwright().start()
        _state["browser"] = _state["pw"].chromium.launch(headless=False)
        _state["page"] = _state["browser"].new_page()
    return _state["page"]


@tool("browser_open", "Open a URL in the controlled browser and return the page text.", {"url": "Full https URL"})
def browser_open(url):
    if not url.startswith(("http://", "https://")):
        return "Only http(s) URLs are allowed."
    def run():
        p = _page(); p.goto(url, timeout=30000)
        return p.inner_text("body")[:3000]
    return _pw.submit(run).result()


@tool("youtube_play", "Search YouTube and play the first matching video.", {"query": "What to play"})
def youtube_play(query):
    def run():
        p = _page()
        p.goto("https://www.youtube.com/results?search_query=" + quote_plus(query), timeout=30000)
        p.click("ytd-video-renderer a#video-title", timeout=15000)
        return "Playing: " + p.title()
    return _pw.submit(run).result()


# ---------- documents ----------
@tool("read_document", "Read an uploaded PDF, DOCX, XLSX or text file by file name.", {"filename": "Name of the uploaded file"})
def read_document(filename):
    path = UPLOADS / Path(filename).name
    if not path.exists():
        return "File not found. Upload it first."
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            text = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
    elif ext == ".docx":
        import docx
        text = "\n".join(p.text for p in docx.Document(path).paragraphs)
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        text = "\n".join(f"[{ws.title}] " + " | ".join(str(c) for c in row if c is not None)
                         for ws in wb for row in ws.iter_rows(values_only=True))
    else:
        text = path.read_text(errors="ignore")
    return text[:12000] or "No text found."


# ---------- screen vision ----------
@tool("look_at_screen", "Take a screenshot of the user's screen and answer a question about it.", {"question": "What to look for"}, confirm=True)
def look_at_screen(question):
    import mss, mss.tools
    with mss.mss() as s:
        shot = s.grab(s.monitors[1])
        png = mss.tools.to_png(shot.rgb, shot.size)
    b64 = base64.b64encode(png).decode()
    out = chat_sync([{"role": "user", "content": [
        {"type": "text", "text": question},
        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}}]}])
    return out["choices"][0]["message"]["content"]


# ---------- messages ----------
@tool("send_email", "Send an email from the user's configured account.", {"to": "Recipient address", "subject": "Subject", "body": "Message body"}, confirm=True)
def send_email(to, subject, body):
    host, user, pw = os.getenv("SMTP_HOST"), os.getenv("SMTP_USER"), os.getenv("SMTP_PASS")
    if not (host and user and pw):
        return "Email is not configured. Set SMTP_HOST, SMTP_USER and SMTP_PASS in .env."
    m = EmailMessage(); m["From"], m["To"], m["Subject"] = user, to, subject; m.set_content(body)
    with smtplib.SMTP_SSL(host, int(os.getenv("SMTP_PORT", "465"))) as s:
        s.login(user, pw); s.send_message(m)
    return "Email sent."


# ---------- LLM access ----------
def chat_sync(messages, tools=None):
    body = {"model": os.getenv("LLM_MODEL", "gpt-4o-mini"), "messages": messages}
    if tools:
        body["tools"] = tools
    key = os.getenv("LLM_API_KEY", "")
    r = httpx.post(os.getenv("LLM_BASE_URL", "https://api.openai.com/v1").rstrip("/") + "/chat/completions",
                   json=body, headers={"Authorization": f"Bearer {key}"} if key else {}, timeout=90)
    r.raise_for_status()
    return r.json()


# ---------- plugins: every plugins/*.py with register(tool) is loaded ----------
def load_plugins():
    import importlib.util
    for f in sorted((ROOT / "plugins").glob("*.py")):
        if f.name.startswith("_"):
            continue
        spec = importlib.util.spec_from_file_location(f"plugin_{f.stem}", f)
        mod = importlib.util.module_from_spec(spec)
        try:
            spec.loader.exec_module(mod)
            mod.register(tool)
        except Exception as e:  # a broken plugin must never stop the app
            print(f"Plugin {f.name} failed to load: {e}")
